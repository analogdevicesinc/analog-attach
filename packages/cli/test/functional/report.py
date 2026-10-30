#!/usr/bin/env python3
"""
The functional tests' report: which requirements were tested, in which ecosystem, and how
it went.

It joins three things:
  - the tests' tags, read from the .bats files: `# bats test_tags=req:CONFIG-04` above a test,
    `# bats file_tags=eco:noos` for a file (or the ecosystem from its folder, noos/ or linux/);
  - bats' JUnit results (`--report-formatter junit`): pass, fail or skipped per test;
  - the requirements (docs/requirements.csv, from cli-spec.xlsx).

and writes, in the output directory:
  - tests.csv         one row per test: its file, name, ecosystem, requirements, status;
  - requirements.csv  one row per requirement: its status in each ecosystem it applies to
                      (pass, fail, skipped, not tested; '-' where it doesn't apply);
  - report.xlsx       both, as sheets, when openpyxl is installed.

It exits 1 when a test names a requirement that doesn't exist, so a renamed requirement
can't leave a stale test behind; a test with no requirement is only listed.

    report.py --junit out/report.xml --tests test/functional --requirements docs/requirements.csv --out out noos linux
"""

import argparse
import csv
import pathlib
import re
import sys
import xml.etree.ElementTree as ElementTree

ECOSYSTEMS = ["linux", "noos"]
# The Ecosystems column of the requirements, as the ecosystems a requirement applies to.
APPLIES = {"All": ECOSYSTEMS, "Linux": ["linux"], "CAPI": ["noos"]}

TAGS_LINE = re.compile(r"^\s*#\s*bats\s+(test_tags|file_tags)\s*=\s*(.*)$")
TEST_LINE = re.compile(r"^\s*@test\s+(['\"])(.*)\1\s*\{")


def split_tags(text):
    return [tag.strip() for tag in text.split(",") if tag.strip()]


def read_tests(root):
    """Every test of the .bats files under root: {(relative file, name): {"eco", "reqs"}}."""
    tests = {}
    for file in sorted(root.rglob("*.bats")):
        relative = file.relative_to(root).as_posix()
        file_tags, pending = [], []
        for line in file.read_text(encoding="utf-8").splitlines():
            tags = TAGS_LINE.match(line)
            if tags:
                (file_tags if tags.group(1) == "file_tags" else pending).extend(split_tags(tags.group(2)))
                continue
            test = TEST_LINE.match(line)
            if test:
                all_tags = file_tags + pending
                pending = []
                eco = next((tag[4:] for tag in all_tags if tag.startswith("eco:")), file.parent.name)
                reqs = [tag[4:] for tag in all_tags if tag.startswith("req:")]
                tests[(relative, test.group(2))] = {"eco": eco, "reqs": reqs}
    return tests


def test_file(classname, root, targets):
    """The .bats file a JUnit classname names, relative to root. bats writes it relative to the
    first folder it was given and absolute for the others, so it is resolved against each."""
    path = pathlib.Path(classname)
    bases = [root, *((target if target.is_dir() else target.parent) for target in targets)]
    for candidate in ([path] if path.is_absolute() else [base / path for base in bases]):
        if candidate.is_file():
            return candidate.resolve().relative_to(root.resolve()).as_posix()
    return classname


def read_results(junit, root, targets):
    """bats' JUnit: {(file relative to root, name): (status, message, seconds)}."""
    results = {}
    for case in ElementTree.parse(junit).getroot().iter("testcase"):
        failure = case.find("failure")
        skipped = case.find("skipped")
        if failure is not None:
            status, message = "fail", (failure.text or "").strip()
        elif skipped is not None:
            status, message = "skipped", (skipped.get("message") or skipped.text or "").strip()
        else:
            status, message = "pass", ""
        results[(test_file(case.get("classname"), root, targets), case.get("name"))] = (status, message, case.get("time", ""))
    return results


def combined(statuses):
    """One status for a requirement in an ecosystem, from its tests' statuses."""
    if not statuses:
        return "not tested"
    if "fail" in statuses:
        return "fail"
    if "pass" in statuses:
        return "pass"
    return "skipped"


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def write_xlsx(path, sheets):
    """Each sheet a (title, header, rows); colour the statuses. Skipped without openpyxl."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return False
    fills = {"pass": "d1fae5", "fail": "fee2e2", "skipped": "fef9c3", "not tested": "f3f4f6"}
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, header, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.value in fills:
                    cell.fill = PatternFill("solid", fgColor=fills[cell.value])
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--junit", required=True, type=pathlib.Path)
    parser.add_argument("--tests", required=True, type=pathlib.Path, help="the folder of the .bats files")
    parser.add_argument("--requirements", required=True, type=pathlib.Path)
    parser.add_argument("--out", required=True, type=pathlib.Path)
    parser.add_argument("targets", nargs="*", type=pathlib.Path, help="the files and folders bats ran, relative to --tests")
    args = parser.parse_args()

    with open(args.requirements, encoding="utf-8") as handle:
        requirements = list(csv.DictReader(handle))
    known = {row["ID"] for row in requirements}
    tests = read_tests(args.tests)
    targets = [target if target.is_absolute() else args.tests / target for target in args.targets]
    results = read_results(args.junit, args.tests, targets) if args.junit.exists() else {}

    test_rows = []
    by_requirement = {}  # (id, eco) -> [status]
    unknown, untagged = [], []
    for (file, name), test in tests.items():
        status, message, seconds = results.get((file, name), ("not run", "", ""))
        test_rows.append([file, name, test["eco"], " ".join(test["reqs"]), status, seconds, message.splitlines()[0] if message else ""])
        if not test["reqs"]:
            untagged.append(f"{file}: {name}")
        for req in test["reqs"]:
            if req not in known:
                unknown.append(f"{file}: {name}: {req}")
            if status != "not run":
                by_requirement.setdefault((req, test["eco"]), []).append(status)

    requirement_rows = []
    for row in requirements:
        applies = APPLIES.get(row["Ecosystems"], ECOSYSTEMS)
        statuses = [combined(by_requirement.get((row["ID"], eco), [])) if eco in applies else "-" for eco in ECOSYSTEMS]
        count = sum(len(by_requirement.get((row["ID"], eco), [])) for eco in ECOSYSTEMS)
        requirement_rows.append([row["ID"], row["Command"], row["Name"], row["Ecosystems"], row["Basis"], *statuses, count])

    args.out.mkdir(parents=True, exist_ok=True)
    test_header = ["File", "Test", "Ecosystem", "Requirements", "Status", "Seconds", "Message"]
    requirement_header = ["ID", "Command", "Name", "Ecosystems", "Basis", *(f"{eco}" for eco in ECOSYSTEMS), "Tests"]
    write_csv(args.out / "tests.csv", test_header, test_rows)
    write_csv(args.out / "requirements.csv", requirement_header, requirement_rows)
    xlsx = write_xlsx(args.out / "report.xlsx", [("Requirements", requirement_header, requirement_rows), ("Tests", test_header, test_rows)])

    statuses = [row[4] for row in test_rows]
    applicable = [status for row in requirement_rows for status in row[5:7] if status != "-"]
    tested = sum(1 for row in requirement_rows if any(status not in ("-", "not tested") for status in row[5:7]))
    print(f"tests: {statuses.count('pass')} passed, {statuses.count('fail')} failed, {statuses.count('skipped')} skipped"
          + (f", {statuses.count('not run')} not run" if "not run" in statuses else ""))
    print(f"requirements: {tested} of {len(requirement_rows)} tested"
          f" ({applicable.count('pass')} pass, {applicable.count('fail')} fail, {applicable.count('skipped')} skipped per ecosystem)")
    print(f"report: {args.out / 'requirements.csv'}, {args.out / 'tests.csv'}" + (f", {args.out / 'report.xlsx'}" if xlsx else ""))
    for line in untagged:
        print(f"warning: no requirement: {line}", file=sys.stderr)
    for line in unknown:
        print(f"error: unknown requirement: {line}", file=sys.stderr)
    return 1 if unknown else 0


if __name__ == "__main__":
    sys.exit(main())
